"""Offline regression tests: strict no-guessing extraction and the 2000+YY two-digit-year policy.

Synthetic ground truth only. No Gemini, no Supabase, no network. The clock is pinned (TODAY) so nothing
depends on the wall clock.
"""
from datetime import datetime

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import process
from app.domain.models import ExtractionState, FieldValue, ReviewState
from app.services.extraction import gemini_adapter
from app.services.extraction.commitment_classifier import classify_commitments
from app.services.extraction.extractor import ExtractionResult
from app.services.extraction.mock_extractor import MockExtractor
from app.services.processing.background import BackgroundProcessor, find_retry_targets, merge_retry_results
from app.services.retry.focused_retry import FocusedRetryOrchestrator
from app.services.storage.json_storage import JSONStorage
from app.services.validation import date_normalizer, validator
from app.services.validation.date_normalizer import TWO_DIGIT_YEAR_CENTURY, assess_date

TODAY = datetime(2026, 6, 1)
BATCH = "batch_policy"
COMMITMENT = "התחייבות"
REQUEST = "דרישה"
REVIEW_FIELDS = ("patient_first_name", "patient_last_name", "patient_id_number",
                 "sender", "received_date", "central_numbers", "commitment_number")


@pytest.fixture(autouse=True)
def pinned_clock(monkeypatch):
    monkeypatch.setattr(date_normalizer, "_now", lambda: TODAY)


# ---------------------------------------------------------------- helpers
def _fields(**overrides):
    base = {
        "patient_first_name": "Dana", "patient_last_name": "Levi", "patient_id_number": "012345674",
        "sender": "Clalit", "received_date": "15/03/2024", "central_numbers": ["946621"],
        "commitment_numbers": ["730578626"], "commitment_types": [{"number": "730578626", "label": COMMITMENT}],
        "commitment_candidates": [],
    }
    base.update(overrides)
    return base


def _result(warnings=(), **overrides):
    return ExtractionResult("1", _fields(**overrides), list(warnings), "synthetic", "test", 0.0)


def _record(warnings=(), **overrides):
    processor = object.__new__(BackgroundProcessor)
    return processor._create_commitment_from_page(_result(warnings, **overrides), "doc", 0)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(process, "_storage", JSONStorage(tmp_path))
    monkeypatch.setattr(process.settings, "temp_storage_path", str(tmp_path))
    app = FastAPI()
    app.include_router(process.router)
    return TestClient(app)


def _store(record):
    process._storage.save_batch(BATCH, {"test_records": [], "commitment_records": [record.model_dump(mode="json")]})


def _verify(client, record, name):
    return client.post(f"/api/batches/{BATCH}/tests/{record.commitment_record_id}/fields/{name}/verify")


def _stored(client):
    return client.get(f"/api/batches/{BATCH}").json()["commitment_records"][0]


def _export_rows(client, tmp_path):
    data = client.get(f"/api/batches/{BATCH}").json()
    for record in data["commitment_records"]:
        approved = client.post(f"/api/batches/{BATCH}/commitments/{record['commitment_record_id']}/confirm-review",
                               json={"reviewed": True, "review_token": record["final_review_token"]})
        assert approved.status_code == 200
    response = client.post(f"/api/batches/{BATCH}/export")
    assert response.status_code == 200
    out = tmp_path / "read_back.xlsx"
    out.write_bytes(response.content)
    return list(load_workbook(out).active.iter_rows(min_row=2, values_only=True))


# ---------------------------------------------------------------- 2000+YY policy (normalizer)
def test_century_constant_is_single_source():
    assert TWO_DIGIT_YEAR_CENTURY == 2000


@pytest.mark.parametrize("text,expected", [
    ("15/03/24", "15/03/2024"),
    ("15.03.24", "15/03/2024"),   # ground truth from the request
    ("01-01-25", "01/01/2025"),   # ground truth from the request
    ("20/04/26", "20/04/2026"),   # DD/MM/26 -> 2026, never 1926
])
def test_two_digit_year_is_read_as_2000_plus_yy_and_flagged(text, expected):
    result = assess_date(text, TODAY)
    assert result.status == "century_inferred" and result.century_inferred is True
    assert result.normalized == expected
    assert not result.normalized.endswith("/1926") and "/19" not in result.normalized[5:]


@pytest.mark.parametrize("text,year", [("15/03/2026", "2026"), ("15/03/1926", "1926"), ("2026-03-15", "2026")])
def test_explicit_four_digit_years_are_preserved_without_inference(text, year):
    result = assess_date(text, TODAY)
    assert result.status == "ok" and result.century_inferred is False
    assert result.normalized.endswith(year)


def test_two_digit_99_is_rejected_as_future_never_1999():
    # 99 -> 2099 is after today: the unchanged future-date policy rejects it (not remapped to the past).
    result = assess_date("01/01/99", TODAY)
    assert result.status == "future" and result.normalized == "01/01/2099"
    assert "1999" not in (result.normalized or "")
    issue = validator.assess_date("01/01/99", today=TODAY)[1]
    assert issue.severity == "error"


@pytest.mark.parametrize("text", ["31/02/26", "29/02/25", "00/01/26", "13/13/26", "31/04/2024"])
def test_impossible_dates_are_rejected_without_a_normalized_value(text):
    result = assess_date(text, TODAY)
    assert result.status == "invalid_date" and result.normalized is None
    value, issue = validator.assess_date(text, today=TODAY)
    assert value is None and issue.severity == "error"


def test_leap_days_and_unrecognized_text():
    assert assess_date("29/02/24", TODAY).normalized == "29/02/2024"
    assert assess_date("29/02/00", TODAY).normalized == "29/02/2000"
    for text in ("ט\"ו באדר", "15/03/24 10:30", "15/3"):
        assert assess_date(text, TODAY).status == "unrecognized"
    assert assess_date(None, TODAY).status == "empty"


def test_future_policy_is_unchanged_for_explicit_years():
    assert assess_date("02/06/2026", TODAY).status == "future"
    assert assess_date("01/06/2026", TODAY).status == "ok"


def test_validator_marks_inferred_century_as_warning_with_code():
    value, issue = validator.assess_date("15/03/24", today=TODAY)
    assert value == "15/03/2024"
    assert issue.severity == "warning" and issue.code == "century_inferred"
    assert issue.message.startswith("Century inferred: ")
    assert validator.assess_date("15/03/2024", today=TODAY) == ("15/03/2024", None)


# ---------------------------------------------------------------- pipeline: no false confidence
def test_inferred_date_keeps_raw_text_and_needs_review():
    date = _record(received_date="15/03/24").received_date
    assert date.raw_value == "15/03/24"
    assert date.normalized_value == "15/03/2024"
    assert date.extraction_state == ExtractionState.uncertain
    assert date.review_state == ReviewState.needs_review
    assert [i.code for i in date.validation_issues] == ["century_inferred"]
    assert date.reviewed_value is None


def test_explicit_four_digit_date_is_not_flagged():
    date = _record(received_date="15/03/2024").received_date
    assert date.extraction_state == ExtractionState.found and not date.validation_issues


def test_invalid_date_is_marked_invalid_and_not_rewritten():
    date = _record(received_date="31/02/26").received_date
    assert date.extraction_state == ExtractionState.invalid
    assert date.normalized_value == "31/02/26"  # untouched text, never replaced by a valid-looking date
    assert date.review_state == ReviewState.needs_review


def test_missing_names_stay_empty_and_flagged():
    for blank in (None, "  ", "לא קריא", "N/A"):
        record = _record(patient_first_name=blank, patient_last_name=blank)
        for field in (record.patient_first_name, record.patient_last_name):
            assert field.raw_value is None and field.normalized_value is None
            assert field.extraction_state == ExtractionState.missing
            assert field.review_state == ReviewState.needs_review


def test_sender_text_is_never_copied_into_patient_name():
    record = _record(patient_first_name=None, patient_last_name=None, sender="Maccabi")
    assert record.patient_first_name.raw_value is None and record.patient_last_name.raw_value is None
    assert record.sender.raw_value == "Maccabi"


def test_invalid_id_is_marked_invalid():
    for bad in ("000000000", "012345675"):
        assert _record(patient_id_number=bad).patient_id_number.extraction_state == ExtractionState.invalid


def test_unreadable_handwriting_warning_flags_field_not_found():
    record = _record(["received_date: handwriting unreadable", "central_numbers: digits unreadable"],
                     received_date=None, central_numbers=[])
    assert record.received_date.raw_value is None
    assert record.received_date.extraction_state != ExtractionState.found
    assert record.received_date.review_state == ReviewState.needs_review
    assert record.central_numbers.raw_value in ([], None)
    assert record.central_numbers.review_state == ReviewState.needs_review


def test_central_numbers_warning_downgrades_a_found_value():
    mapped = process._map_extraction_to_test_record(_result(["central_numbers: second digit unclear"]), "doc", 0)
    process._detect_uncertainty(mapped, _result(["central_numbers: second digit unclear"]))
    assert mapped.central_numbers.extraction_state == ExtractionState.uncertain
    assert mapped.central_numbers.review_state == ReviewState.needs_review


def test_central_number_list_is_cleaned_but_keeps_leading_zeros_and_every_number():
    mapped = process._map_extraction_to_test_record(
        _result(central_numbers=["00123", " 946621 ", "", "null", "946621", "00124"]), "doc", 0)
    assert mapped.central_numbers.raw_value == ["00123", "946621", "00124"]


def test_fully_empty_extraction_leaves_every_field_empty_and_blocks_export():
    record = _record(patient_first_name=None, patient_last_name=None, patient_id_number=None, sender=None,
                     received_date=None, central_numbers=[], commitment_numbers=[], commitment_types=[])
    for name in ("patient_first_name", "patient_last_name", "patient_id_number", "sender", "received_date"):
        field = getattr(record, name)
        assert field.raw_value is None and field.extraction_state == ExtractionState.missing
    can_export, blockers = process._check_export_eligibility_commitments([record.model_dump(mode="json")])
    assert not can_export
    assert any(b.endswith("patient_id_number_missing") for b in blockers)
    assert any(b.endswith("commitment_number_missing") for b in blockers)


def test_conflicting_scalar_values_across_pages_are_flagged_not_silently_chosen():
    page0 = _result(patient_id_number="012345674")
    page1 = _result(patient_id_number="123456782")
    merged_record, merged = process._merge_multi_page_extractions(page0, page1, "doc", 0, 1)
    assert any(w.startswith("patient_id_number: conflicting values across pages") for w in merged.warnings)
    process._detect_uncertainty(merged_record, merged)
    assert merged_record.patient_id_number.review_state == ReviewState.needs_review
    assert merged_record.patient_id_number.extraction_state != ExtractionState.found


# ---------------------------------------------------------------- multiple / ambiguous numbers
def _multi_page():
    numbers = ["730578626", "730578627", "730578628"]
    types = [{"number": numbers[0], "label": COMMITMENT}, {"number": numbers[1], "label": REQUEST},
             {"number": numbers[2], "label": COMMITMENT}, {"number": "555", "label": "בדיקה"}]
    central = ["0946621", "0946622"]
    classification = classify_commitments(numbers + ["555"], types, central, "Clalit")
    return _record(central_numbers=central, commitment_numbers=classification.accepted,
                   commitment_types=types, commitment_candidates=classification.candidates)


def test_multiple_numbers_are_preserved_flagged_and_exported_only_after_human_verification(client, tmp_path):
    record = _multi_page()
    field = record.commitment_number
    # Exception-only review rule: several cleanly accepted numbers are not ambiguity, but the unlabeled "555"
    # stays a pending candidate, so the page is still blocked until a human resolves it.
    assert field.extraction_state == ExtractionState.found
    assert {c.value for c in field.candidates} == {"730578626", "730578627", "730578628", "555"}
    _store(record)
    assert client.get(f"/api/batches/{BATCH}/export-check").json()["can_export"] is False
    for name in REVIEW_FIELDS:
        assert _verify(client, record, name).status_code == 200
    rows = _export_rows(client, tmp_path)
    assert [r[6] for r in rows] == ["730578626", "730578627", "730578628"]  # ambiguous "555" not exported
    assert all(r[5] == "0946621, 0946622" for r in rows)  # leading zeros preserved


# ---------------------------------------------------------------- human corrections survive automation
def test_retry_failed_never_replaces_reviewed_records():
    record = _record(received_date="15/03/24")
    stored = record.model_dump(mode="json")
    stored["received_date"]["reviewed_value"] = "15/03/2024"
    stored["received_date"]["review_state"] = "human_verified"
    batch = {"documents": [{"document_id": "doc", "page_count": 1}], "commitment_records": [stored]}
    assert find_retry_targets(batch) == set()
    other = _record(received_date="01/01/2020").model_dump(mode="json")
    assert merge_retry_results([stored], [other], {("doc", 0)}) == [stored]


def test_start_processing_conflicts_when_records_exist(client):
    _store(_record())
    assert client.post(f"/api/batches/{BATCH}/process").status_code == 409


def test_focused_retry_never_touches_human_verified_field():
    orchestrator = FocusedRetryOrchestrator(MockExtractor())
    field = FieldValue[str](raw_value="15/03/24", reviewed_value="15/03/2024", review_state=ReviewState.human_verified,
                            extraction_state=ExtractionState.found)
    orchestrator.update_field_with_retry_result("received_date", field, _result(received_date="20/04/2020"))
    assert (field.raw_value, field.reviewed_value) == ("15/03/24", "15/03/2024")
    assert field.review_state == ReviewState.human_verified
    assert getattr(field, "retry_count", 0) == 0
    assert field.candidates == []


def test_focused_retry_adds_candidate_instead_of_replacing_a_value():
    orchestrator = FocusedRetryOrchestrator(MockExtractor())
    field = FieldValue[str](raw_value="Dana", extraction_state=ExtractionState.uncertain)
    orchestrator.update_field_with_retry_result("patient_first_name", field, _result(patient_first_name="Dina"))
    assert field.raw_value == "Dana"
    assert [(c.value, c.reason) for c in field.candidates] == [("Dina", "retry_alternative")]
    assert field.extraction_state == ExtractionState.uncertain and field.review_state == ReviewState.needs_review


def test_focused_retry_fills_only_empty_fields_and_never_as_found():
    orchestrator = FocusedRetryOrchestrator(MockExtractor())
    field = FieldValue[str](raw_value=None, extraction_state=ExtractionState.missing)
    orchestrator.update_field_with_retry_result("patient_first_name", field, _result(patient_first_name="Dana"))
    assert field.raw_value == "Dana"
    assert field.extraction_state == ExtractionState.uncertain and field.review_state == ReviewState.needs_review
    assert field.retry_count == 1


# ---------------------------------------------------------------- verify / update flow and export blocking
def _inferred_record():
    return _record(received_date="15/03/24")


def test_unverified_inferred_date_blocks_export_until_human_verifies(client, tmp_path):
    record = _inferred_record()
    _store(record)
    check = client.get(f"/api/batches/{BATCH}/export-check").json()
    assert check["can_export"] is False
    assert any(b.endswith("received_date_unreviewed") for b in check["blockers"])
    for name in REVIEW_FIELDS:
        if name != "received_date":
            assert _verify(client, record, name).status_code == 200
    assert client.get(f"/api/batches/{BATCH}/export-check").json()["can_export"] is False
    assert client.post(f"/api/batches/{BATCH}/export").status_code != 200

    approved = _verify(client, record, "received_date").json()["received_date"]
    assert approved["reviewed_value"] == "15/03/2024"
    assert approved["review_state"] == "human_verified" and approved["extraction_state"] == "found"
    assert approved["raw_value"] == "15/03/24"  # original text preserved
    assert not [i for i in approved["validation_issues"] if i.get("code") == "century_inferred"]
    rows = _export_rows(client, tmp_path)
    assert "15/03/2024" in rows[0] and "15/03/24" not in rows[0]


def test_second_verify_does_not_revert_the_stored_value(client):
    record = _inferred_record()
    _store(record)
    _verify(client, record, "received_date")
    again = _verify(client, record, "received_date").json()["received_date"]
    assert again["reviewed_value"] == "15/03/2024" and again["review_state"] == "human_verified"


@pytest.mark.parametrize("typed,stored", [("15/03/24", "15/03/2024"), ("01-01-25", "01/01/2025"),
                                          ("15/03/1926", "15/03/1926"), ("15/03/2026", "15/03/2026")])
def test_update_field_accepts_valid_dates_and_stores_four_digit_value(client, typed, stored):
    record = _inferred_record()
    _store(record)
    path = f"/api/batches/{BATCH}/tests/{record.commitment_record_id}/fields/received_date"
    response = client.patch(path, json={"new_value": typed})
    assert response.status_code == 200
    field = _stored(client)["received_date"]
    assert field["reviewed_value"] == stored and field["review_state"] == "human_verified"
    assert not [i for i in field["validation_issues"] if i.get("code") == "century_inferred"]


@pytest.mark.parametrize("typed", ["01/01/99", "31/02/26", "13/13/26", "abc"])
def test_update_field_rejects_impossible_or_future_inferred_dates(client, typed):
    record = _inferred_record()
    _store(record)
    path = f"/api/batches/{BATCH}/tests/{record.commitment_record_id}/fields/received_date"
    assert client.patch(path, json={"new_value": typed}).status_code == 400
    assert _stored(client)["received_date"]["reviewed_value"] is None


def test_verify_rejects_invalid_date(client):
    record = _record(received_date="31/02/26")
    _store(record)
    assert _verify(client, record, "received_date").status_code == 400
    assert _stored(client)["received_date"]["review_state"] != "human_verified"


def test_validation_rerun_does_not_overwrite_a_reviewed_date():
    mapped = process._map_extraction_to_test_record(_result(received_date="15/03/24"), "doc", 0)
    mapped.received_date.reviewed_value = "15/03/2023"
    mapped.received_date.normalized_value = "15/03/2023"
    mapped.received_date.review_state = ReviewState.human_verified
    process._validate_record(mapped)
    assert mapped.received_date.reviewed_value == "15/03/2023"
    assert mapped.received_date.normalized_value == "15/03/2023"
    assert mapped.received_date.review_state == ReviewState.human_verified
    assert not mapped.received_date.validation_issues


# ---------------------------------------------------------------- prompt safeguards (string level, no network)
PROMPT = gemini_adapter.EXTRACTION_PROMPT_V2


@pytest.mark.parametrize("needle", [
    "מרכזי", "מספר מרכזי", "התחייבות", "מספר התחייבות", "דרישה", "דגימה", "התקבל",
])
def test_prompt_contains_hebrew_labels(needle):
    assert needle in PROMPT


@pytest.mark.parametrize("needle", [
    "extract only what is visibly supported",
    "NEVER complete, extend, pad or reconstruct a partially visible number",
    "NEVER take a patient name from the sender",
    "a signature, a stamp",
    "NEVER derive a first or last name from the other name",
    "Preserve EVERY identifiable central number: no cap",
    "Preserve EVERY identifiable commitment number: no cap",
    "leading zeros kept as strings",
    "A nearby number without evidence of its meaning",
    "A nearby number is NEVER a commitment number without its label",
    "do NOT silently choose one",
    "naming the field and the alternatives",
    "Never mark an uncertain field as confident",
    "a two-digit year stays two-digit (never expand it, never guess the century)",
    "must begin with the exact snake_case field name",
])
def test_prompt_contains_no_guessing_safeguards(needle):
    assert needle.lower() in PROMPT.lower()


def test_prompt_names_every_schema_field_in_warning_format_and_output_block():
    for name in gemini_adapter.CommitmentFormExtraction.model_fields:
        assert f'"{name}"' in PROMPT.split("OUTPUT FORMAT:")[1]
    for name in ("patient_first_name", "patient_last_name", "patient_id_number", "sender", "received_date",
                 "central_numbers", "commitment_numbers"):
        assert name in PROMPT.split("VALIDATION AND WARNINGS:")[1].split("ABSOLUTE RULES:")[0]


def test_response_schema_fields_are_unchanged():
    assert list(gemini_adapter.CommitmentFormExtraction.model_fields) == [
        "patient_first_name", "patient_last_name", "patient_id_number", "sender", "received_date",
        "central_numbers", "commitment_numbers", "commitment_types", "warnings"]
    assert gemini_adapter.PROMPT_VERSION_V2 == "2.0"
