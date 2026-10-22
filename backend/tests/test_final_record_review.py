"""Offline API/export tests for explicit, data-bound whole-record review."""
from io import BytesIO
import pytest
from fastapi import FastAPI, Depends
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.api import process
from app.core.access import require_access
from app.domain.models import CommitmentRecord, FieldValue
from app.services.review_rules import annotate_record, final_review_confirmed
from app.services.storage.json_storage import JSONStorage

BATCH = 'batch_final_review'
RECORD = 'record_review'


def clean_record():
    values = {'patient_first_name':'Dana','patient_last_name':'Levi','patient_id_number':'012345674',
              'sender':'Clalit','received_date':'15/03/2024','central_numbers':['00121','00122'],
              'commitment_number':'00031'}
    fields = {k:FieldValue(raw_value=v,normalized_value=v,extraction_state='found',review_state='auto_accepted')
              for k,v in values.items()}
    return CommitmentRecord(commitment_record_id=RECORD,form_id='doc_page_0',source_page=1,
                            relationship_status='inferred',**fields).model_dump(mode='json')


@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(process,'_storage',JSONStorage(tmp_path))
    monkeypatch.setattr(process.settings,'temp_storage_path',str(tmp_path))
    app=FastAPI()
    app.include_router(process.router,dependencies=[Depends(require_access)])
    process._storage.save_batch(BATCH,{'status':'completed','test_records':[],
                                      'commitment_records':[clean_record()],'documents':[]})
    return TestClient(app)


def record(client):
    return client.get(f'/api/batches/{BATCH}').json()['commitment_records'][0]


def approve(client,data=None):
    data=data or record(client)
    return client.post(f'/api/batches/{BATCH}/commitments/{RECORD}/confirm-review',
                       json={'reviewed':True,'review_token':data['final_review_token']})


def test_clean_ai_values_blocked_until_explicit_record_confirmation(client):
    data=record(client)
    assert not data['final_review_confirmed']
    check=client.get(f'/api/batches/{BATCH}/export-check').json()
    assert not check['can_export']
    assert f'{RECORD}: final_record_review_required' in check['blockers']
    assert client.post(f'/api/batches/{BATCH}/export').status_code==400
    assert approve(client).status_code==200
    assert all(record(client)[name]['review_state']=='auto_accepted' for name in (
        'patient_first_name','patient_last_name','patient_id_number','sender','received_date',
        'central_numbers','commitment_number'))
    response=client.post(f'/api/batches/{BATCH}/export')
    assert response.status_code==200
    rows=list(load_workbook(BytesIO(response.content)).active.values)
    assert any('012345674' in row and '00031' in row for row in rows[1:])


def test_approval_persists_and_stale_display_cannot_confirm_new_values(client):
    old=record(client)
    assert approve(client).status_code==200
    saved=process._storage.load_batch(BATCH)['commitment_records'][0]
    assert final_review_confirmed(CommitmentRecord(**saved).model_dump(mode='json'))
    timestamp=saved['final_review_confirmed_at']
    assert approve(client).json()['final_review_confirmed_at']==timestamp
    response=client.patch(f'/api/batches/{BATCH}/tests/{RECORD}/fields/patient_first_name',json={'new_value':'Dina'})
    assert response.status_code==200  # auto_accepted remains editable
    assert not record(client)['final_review_confirmed']
    assert process._storage.load_batch(BATCH)['commitment_records'][0]['final_review_fingerprint'] is None
    assert approve(client,old).status_code==409
    assert approve(client).status_code==200
    # Saving identical confirmed data does not demand another whole-record approval.
    client.patch(f'/api/batches/{BATCH}/tests/{RECORD}/fields/patient_first_name',json={'new_value':'Dina'})
    assert record(client)['final_review_confirmed']


@pytest.mark.parametrize('change', ['missing','invalid','ambiguous','relationship'])
def test_record_approval_cannot_clear_unresolved_blockers(client,change):
    batch=process._storage.load_batch(BATCH)
    item=batch['commitment_records'][0]
    if change=='relationship':
        item['relationship_status']='unresolved_relationship'
    else:
        field=item['patient_id_number']
        field['review_state']='needs_review'
        field['extraction_state']={'missing':'missing','invalid':'invalid','ambiguous':'conflicting'}[change]
        if change=='missing':
            field['raw_value']=field['normalized_value']=None
        if change=='invalid':
            field['validation_issues']=[{'field_name':'patient_id_number','severity':'error','message':'checksum'}]
    process._storage.save_batch(BATCH,batch)
    assert approve(client).status_code==400
    assert not record(client)['final_review_confirmed']
    assert client.post(f'/api/batches/{BATCH}/export').status_code==400


def test_multi_numbers_human_values_and_confirmed_empty_survive(client):
    batch=process._storage.load_batch(BATCH)
    item=batch['commitment_records'][0]
    item['patient_first_name'].update(reviewed_value='',review_state='human_verified')
    item['patient_last_name'].update(reviewed_value='Reviewed',review_state='human_verified')
    process._storage.save_batch(BATCH,batch)
    response=client.patch(f'/api/batches/{BATCH}/tests/{RECORD}/fields/commitment_number',
                          json={'new_value':['00031','00032']})
    assert response.status_code==200
    # Changing which number is primary must not discard the previous primary when explicitly retained.
    client.patch(f'/api/batches/{BATCH}/tests/{RECORD}/fields/commitment_number',json={'new_value':['00032','00031']})
    assert approve(client).status_code==200
    item=record(client)
    client.patch(f'/api/batches/{BATCH}/tests/{RECORD}/fields/commitment_number',json={'new_value':['00032','00031']})
    assert record(client)['final_review_confirmed']  # unchanged numbers do not demand repeated approval
    assert item['patient_first_name']['reviewed_value']==''
    assert item['patient_last_name']['reviewed_value']=='Reviewed'
    assert item['central_numbers']['raw_value']==['00121','00122']
    response=client.post(f'/api/batches/{BATCH}/export')
    rows=list(load_workbook(BytesIO(response.content)).active.values)[1:]
    assert len(rows)==2
    assert {row[-1] for row in rows}=={'00031','00032'}


def test_data_change_from_extraction_and_candidate_change_invalidate_even_if_approval_copied(client):
    approve(client)
    batch=process._storage.load_batch(BATCH)
    item=batch['commitment_records'][0]
    item['central_numbers']['raw_value']=item['central_numbers']['normalized_value']=['00999']
    # Simulate a different writer retaining old confirmation metadata; export still checks the current data.
    process._storage.save_batch(BATCH,batch)
    assert not record(client)['final_review_confirmed']
    assert client.post(f'/api/batches/{BATCH}/export').status_code==400
    approve(client)
    batch=process._storage.load_batch(BATCH)
    batch['commitment_records'][0]['commitment_number']['candidates'].append(
        {'value':'00032','reason':'multiple_commitments_on_page','status':'pending'})
    process._storage.save_batch(BATCH,batch)
    assert not record(client)['final_review_confirmed']


def test_explicit_action_authentication_and_batch_scope(client,monkeypatch):
    url=f'/api/batches/{BATCH}/commitments/{RECORD}/confirm-review'
    assert client.post(url,json={'review_token':record(client)['final_review_token']}).status_code==400
    assert client.post(url.replace(RECORD,'other_record'),json={'reviewed':True}).status_code==404
    monkeypatch.setattr(process.settings,'api_username','secretary')
    monkeypatch.setattr(process.settings,'api_password','synthetic-password-only')
    assert client.post(url,json={'reviewed':True}).status_code==401
