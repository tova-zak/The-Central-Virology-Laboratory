"""Synthetic, offline regressions for the pipeline audit on the final-review baseline."""
import asyncio
from io import BytesIO
from unittest.mock import Mock

import cv2
import numpy as np
import pytest
from fastapi import Depends, FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app.api import process
from app.core.access import require_access
from app.core.config import Settings
from app.domain.models import FieldValue
from app.services.extraction.extractor import ExtractionResult
from app.services.processing import background as bg
from app.services.preprocessing.image_preprocessor import ImagePreprocessor, PreprocessingConfig, PreprocessingMethod
from app.services.storage.json_storage import JSONStorage
from app.services.review_rules import annotate_record, record_review_token


def result(warnings=()):
    return ExtractionResult('1', dict(patient_first_name='Synthetic', patient_last_name='Example',
        patient_id_number='000000018', sender='Clalit', received_date='15/03/2024',
        central_numbers=['00121', '00122'], commitment_numbers=['00031']), list(warnings), 'synthetic', 'test', 0)


def record(warnings=()):
    return object.__new__(bg.BackgroundProcessor)._create_commitment_from_page(result(warnings), 'doc', 0).model_dump(mode='json')


@pytest.fixture
def api(tmp_path, monkeypatch):
    monkeypatch.setattr(process, '_storage', JSONStorage(tmp_path))
    monkeypatch.setattr(process.settings, 'temp_storage_path', str(tmp_path))
    app = FastAPI()
    app.include_router(process.router, dependencies=[Depends(require_access)])
    return TestClient(app)


def save(item):
    process._storage.save_batch('audit', {'batch_id':'audit', 'status':'completed',
        'test_records':[], 'commitment_records':[item], 'documents':[]})


def confirm(api, item):
    return api.post(f"/api/batches/audit/commitments/{item['commitment_record_id']}/confirm-review",
                    json={'reviewed':True, 'review_token':record_review_token(item)})


def png(array):
    return cv2.imencode('.png', array)[1].tobytes()


def aligned():
    a = np.full((600, 800, 3), 255, np.uint8)
    for y in range(80, 500, 50):
        cv2.putText(a, 'SYNTHETIC 00121', (70, y), cv2.FONT_HERSHEY_SIMPLEX, 1, (0,0,0), 2)
    return a


def test_ocr_parser_english_labels_and_id_boundary():
    from app.services.ocr.tesseract_extractor import TesseractExtractor

    text = "\n".join([
        "ID: 0011223344",
        "Date: 2026-06-18",
        "Central: 00121",
        "Commitment: AB1234",
        "noise 31 and 606",
    ])
    fields = TesseractExtractor._parse_ocr_text(TesseractExtractor, text, None)
    assert fields["patient_id_number"] == "0011223344"
    assert fields["received_date"] == "2026-06-18"
    assert fields["central_numbers"] == ["00121"]
    assert fields["commitment_numbers"] == ["AB1234"]


def test_quality_selection_corrects_skew_but_rejects_blur_crop_and_blank():
    good = aligned()
    skewed = cv2.warpAffine(good, cv2.getRotationMatrix2D((400,300), 7, 1), (800,600), borderValue=(255,255,255))
    p = ImagePreprocessor()
    selected, method = p.select_variant(png(skewed), [(png(good), PreprocessingMethod.DESKEWED)])
    assert selected == png(good) and method == PreprocessingMethod.DESKEWED
    for bad in [cv2.GaussianBlur(good, (21,21), 7), good[:150], np.full_like(good,255)]:
        assert p.select_variant(png(good), [(png(bad), PreprocessingMethod.DESKEWED)])[0] == png(good)
    assert p.select_variant(png(good), [(b'broken', PreprocessingMethod.DESKEWED)])[0] == png(good)


def test_original_policy_and_equal_quality_keep_original():
    original = png(aligned())
    p = ImagePreprocessor(PreprocessingConfig(selection_policy='original'))
    assert p.select_variant(original, [(original, PreprocessingMethod.DESKEWED)])[1] == PreprocessingMethod.ORIGINAL
    assert ImagePreprocessor().select_variant(original, [(original, PreprocessingMethod.DESKEWED)])[1] == PreprocessingMethod.ORIGINAL


def test_active_batch_runs_ocr_on_original_bytes(tmp_path, monkeypatch):
    storage = JSONStorage(tmp_path)
    batch = tmp_path / 'audit'
    batch.mkdir()
    original = png(aligned())
    (batch / 'doc.png').write_bytes(original)
    monkeypatch.setattr(bg, 'get_storage_backend', lambda: storage)
    monkeypatch.setattr(bg.settings, 'demo_mode', True)
    monkeypatch.setattr(bg.settings, 'enable_ocr', True)
    monkeypatch.setattr(bg.settings, 'temp_storage_path', str(tmp_path))
    processor = bg.BackgroundProcessor()
    monkeypatch.setattr(processor._hybrid_config, 'enable_ocr', True)
    monkeypatch.setattr(processor._hybrid_config, 'enable_hybrid_verification', False)
    monkeypatch.setattr(processor._hybrid_config, 'enable_preprocessing', False)
    ocr_calls = []

    class FakeOCR:
        model_id = 'tesseract-ocr'

        def extract(self, image, document_id, page_index):
            ocr_calls.append(image)
            return result()

    monkeypatch.setattr(processor, '_get_ocr_extractor', lambda: FakeOCR())
    async def extract(**kwargs):
        return result()
    monkeypatch.setattr(processor, '_extract_with_retry', extract)
    asyncio.run(processor._process_batch(bg.ProcessingJob('audit')))
    assert ocr_calls == [original]


def test_active_batch_sends_selected_bytes_and_persists_method(tmp_path, monkeypatch):
    storage = JSONStorage(tmp_path)
    batch = tmp_path / 'audit'
    batch.mkdir()
    original, selected = png(aligned()), png(cv2.rotate(aligned(), cv2.ROTATE_180))
    (batch/'doc.png').write_bytes(original)
    monkeypatch.setattr(bg, 'get_storage_backend', lambda: storage)
    monkeypatch.setattr(bg.settings, 'demo_mode', True)
    monkeypatch.setattr(bg.settings, 'temp_storage_path', str(tmp_path))
    processor = bg.BackgroundProcessor()
    preprocessor = Mock()
    preprocessor.preprocess.return_value = [(original, PreprocessingMethod.ORIGINAL), (selected, PreprocessingMethod.DESKEWED)]
    preprocessor.select_variant.return_value = (selected, PreprocessingMethod.DESKEWED)
    monkeypatch.setattr(processor, '_get_preprocessor', lambda: preprocessor)
    monkeypatch.setattr(processor._hybrid_config, 'enable_preprocessing', True)
    seen = []
    async def extract(**kwargs):
        seen.append(kwargs['image_bytes'])
        return result()
    monkeypatch.setattr(processor, '_extract_with_retry', extract)
    asyncio.run(processor._process_batch(bg.ProcessingJob('audit')))
    assert seen == [selected]
    saved = storage.load_batch('audit')['commitment_records'][0]
    assert saved['central_numbers']['preprocessing_method'] == 'deskewed'


def test_every_warning_survives_mapping_storage_and_api_and_blocks_review(api):
    warnings = [f'central_numbers: synthetic uncertainty {i} ' + 'x'*350 for i in range(8)]
    warnings += [f'Unscoped uncertainty {i}' for i in range(12)] + ['commitment_numbers: uncertain digits']
    item = record(warnings)
    save(item)
    data = api.get('/api/batches/audit').json()['commitment_records'][0]
    for warning in warnings:
        assert warning in (data['central_numbers']['extraction_notes'] +
                           data['commitment_number']['extraction_notes'] + data['page_warnings'])
    for name in ['central_numbers', 'commitment_number']:
        assert data[name]['review_required']
        assert data[name]['review_state'] != 'human_verified'
    assert confirm(api, item).status_code == 400
    assert api.post('/api/batches/audit/export').status_code == 400


@pytest.mark.parametrize('name', ['patient_first_name', 'patient_last_name', 'sender', 'received_date', 'commitment_number'])
def test_confirmed_empty_exports_empty_and_reverify_preserves_decision(api, name):
    item = record()
    save(item)
    rid = item['commitment_record_id']
    url = f'/api/batches/audit/tests/{rid}/fields/{name}'
    assert api.patch(url, json={'new_value':''}).status_code == 200
    assert api.post(url+'/verify').status_code == 200
    saved = process._storage.load_batch('audit')['commitment_records'][0]
    assert saved[name]['reviewed_value'] == ''
    assert saved[name]['review_state'] == 'human_verified'
    assert confirm(api, saved).status_code == 200
    response = api.post('/api/batches/audit/export')
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content))
    column = ['patient_first_name','patient_last_name','patient_id_number','sender','received_date','central_numbers','commitment_number'].index(name)+1
    assert wb.active.cell(2,column).value is None
    assert wb.active.cell(2,3).value == '000000018'
    wb.close()


@pytest.mark.parametrize('value', [None, [], ''])
def test_null_is_not_an_explicit_confirmed_empty(api, value):
    item = record()
    save(item)
    url = f"/api/batches/audit/tests/{item['commitment_record_id']}/fields/patient_first_name"
    response = api.patch(url, json={'new_value':value})
    assert response.status_code == (200 if value == '' else 400)


@pytest.mark.parametrize('name,value', [('patient_id_number',''), ('central_numbers',[]), ('central_numbers',[''])])
def test_required_empty_cannot_be_confirmed_or_exported(api, name, value):
    item = record()
    save(item)
    url = f"/api/batches/audit/tests/{item['commitment_record_id']}/fields/{name}"
    assert api.patch(url, json={'new_value':value}).status_code == 400


@pytest.mark.parametrize('status,legacy', [('unresolved_relationship','unknown'), ('inferred','ambiguous'), ('unknown','unknown')])
def test_ambiguous_relationship_blocks_direct_export_even_with_matching_record_approval(api, status, legacy):
    item = record()
    item.update(relationship_status=status, relation_status=legacy)
    # A stored whole-record approval is not an override for unresolved relationships.
    item.update(final_review_fingerprint=record_review_token(item), final_review_confirmed_at='synthetic-time')
    save(item)
    assert not api.get('/api/batches/audit/export-check').json()['can_export']
    assert api.post('/api/batches/audit/export').status_code == 400
    assert confirm(api, item).status_code == 400


def test_resolve_validates_refs_updates_both_statuses_and_invalidates_approval(api):
    item = record()
    item.update(relationship_status='unresolved_relationship', relation_status='ambiguous')
    save(item)
    url = f"/api/batches/audit/commitments/{item['commitment_record_id']}/resolve"
    for refs in [[], ['other-document-number'], [121], '00121']:
        assert api.post(url, json={'relation_status':'related','central_number_refs':refs}).status_code == 400
    response = api.post(url, json={'relation_status':'related','central_number_refs':['00121','00122']})
    assert response.status_code == 200
    saved = process._storage.load_batch('audit')['commitment_records'][0]
    assert saved['relationship_status'] == 'explicit'
    assert saved['relation_status'] == 'related'
    assert not api.get('/api/batches/audit/export-check').json()['can_export']  # separate final review still required
    assert confirm(api, saved).status_code == 200
    assert api.post('/api/batches/audit/export').status_code == 200


def test_environment_precedence_and_real_vs_mock_construction(tmp_path, monkeypatch):
    envfile = tmp_path/'synthetic.env'
    envfile.write_text('DEMO_MODE=true\n', encoding='utf-8')
    monkeypatch.delenv('DEMO_MODE', raising=False)
    assert Settings(_env_file=envfile).demo_mode is True
    monkeypatch.setenv('DEMO_MODE','false')
    assert Settings(_env_file=envfile).demo_mode is False
    assert Settings(_env_file=envfile, demo_mode=True).demo_mode is True
    monkeypatch.setattr(bg, 'get_storage_backend', lambda: Mock())
    constructor = Mock()
    monkeypatch.setattr(bg, 'GeminiExtractor', constructor)
    monkeypatch.setattr(bg.settings, 'demo_mode', True)
    assert bg.BackgroundProcessor().extractor_name == 'MockExtractor'
    constructor.assert_not_called()
    monkeypatch.setattr(bg.settings, 'demo_mode', False)
    assert bg.BackgroundProcessor()._extractor is constructor.return_value
    constructor.assert_called_once_with(model=bg.settings.gemini_model)


def test_low_confidence_ocr_cannot_create_conflict_or_overwrite_human_empty():
    from app.services.hybrid.verifier import HybridVerifier, VerificationStatus
    vision, ocr = result(), result(['OCR_LOW_CONFIDENCE: Average confidence 1.0% below threshold 70.0%'])
    ocr.fields['patient_id_number'] = '000000026'
    verifier = HybridVerifier()
    field = FieldValue[str](raw_value='000000018')
    assert verifier.verify_field('patient_id_number', vision, ocr, field).verification_status == VerificationStatus.GEMINI_ONLY
    human = FieldValue[str](raw_value='Old synthetic', reviewed_value='', review_state='human_verified')
    before = human.model_dump()
    evidence = verifier.verify_field('patient_first_name', vision, ocr, human)
    verifier.apply_verification_to_field(human, evidence, 'patient_first_name')
    assert human.model_dump() == before
