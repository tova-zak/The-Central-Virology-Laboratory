"""Tests for Excel export functionality.

TASK-020/TASK-022: Regression and export tests.
"""
import pytest
from pathlib import Path
from openpyxl import load_workbook

from app.services.export import ExcelExporter
from app.domain.models import (
    TestRecord, CommitmentRecord, FieldValue,
    ExtractionState, ReviewState
)


@pytest.fixture
def sample_test_record():
    """Create a sample test record for testing."""
    return TestRecord(
        test_id="test_001",
        document_id="doc_001",
        page_indices=[0],
        patient_first_name=FieldValue[str](
            raw_value="יוסי",
            normalized_value="יוסי",
            extraction_state=ExtractionState.found
        ),
        patient_last_name=FieldValue[str](
            raw_value="כהן",
            normalized_value="כהן",
            extraction_state=ExtractionState.found
        ),
        patient_id_number=FieldValue[str](
            raw_value="123456782",
            normalized_value="123456782",
            extraction_state=ExtractionState.found
        ),
        sender=FieldValue[str](
            raw_value="בית חולים שיבא",
            normalized_value="בית חולים שיבא - תל השומר",
            extraction_state=ExtractionState.found
        ),
        received_date=FieldValue[str](
            raw_value="2024-01-15",
            normalized_value="2024-01-15",
            extraction_state=ExtractionState.found
        ),
        central_numbers=FieldValue[list[str]](
            raw_value=["12345", "67890"],
            normalized_value=["12345", "67890"],
            extraction_state=ExtractionState.found
        ),
        commitment_ids=["comm_001", "comm_002"]
    )


@pytest.fixture
def sample_commitment_records():
    """Create sample commitment records."""
    return [
        CommitmentRecord(
            commitment_record_id="comm_001",
            test_id="test_001",
            commitment_number=FieldValue[str](
                raw_value="ABC123",
                normalized_value="ABC123",
                extraction_state=ExtractionState.found
            ),
            central_number_refs=["12345"],
            relation_status="clear"
        ),
        CommitmentRecord(
            commitment_record_id="comm_002",
            test_id="test_001",
            commitment_number=FieldValue[str](
                raw_value="DEF456",
                normalized_value="DEF456",
                extraction_state=ExtractionState.found
            ),
            central_number_refs=["67890"],
            relation_status="clear"
        )
    ]


def test_excel_exporter_creates_file(tmp_path, sample_test_record, sample_commitment_records):
    """Test that Excel exporter creates a valid file."""
    exporter = ExcelExporter()
    output_path = tmp_path / "test_export.xlsx"
    
    result_path = exporter.export(
        test_records=[sample_test_record],
        commitment_records=sample_commitment_records,
        output_path=output_path
    )
    
    assert result_path.exists()
    assert result_path == output_path


def test_excel_exporter_correct_headers(tmp_path, sample_test_record, sample_commitment_records):
    """Test that Excel file has correct headers."""
    exporter = ExcelExporter()
    output_path = tmp_path / "test_export.xlsx"
    
    exporter.export(
        test_records=[sample_test_record],
        commitment_records=sample_commitment_records,
        output_path=output_path
    )
    
    # Load and check headers
    wb = load_workbook(output_path)
    ws = wb.active
    
    expected_headers = [
        "First Name", "Last Name", "ID Number", "Sender",
        "Received Date", "Central Numbers", "Commitment Numbers"
    ]
    
    for col_idx, expected_header in enumerate(expected_headers, start=1):
        actual_header = ws.cell(row=1, column=col_idx).value
        assert actual_header == expected_header


def test_excel_exporter_correct_data(tmp_path, sample_test_record, sample_commitment_records):
    """Test that Excel file contains correct data."""
    exporter = ExcelExporter()
    output_path = tmp_path / "test_export.xlsx"
    
    exporter.export(
        test_records=[sample_test_record],
        commitment_records=sample_commitment_records,
        output_path=output_path
    )
    
    # Load and check data
    wb = load_workbook(output_path)
    ws = wb.active
    
    # Check row 2 (first data row)
    assert ws.cell(row=2, column=1).value == "יוסי"  # First Name
    assert ws.cell(row=2, column=2).value == "כהן"  # Last Name
    assert ws.cell(row=2, column=3).value == "123456782"  # ID Number
    assert ws.cell(row=2, column=4).value == "בית חולים שיבא - תל השומר"  # Sender (normalized)
    assert ws.cell(row=2, column=5).value == "2024-01-15"  # Received Date
    assert ws.cell(row=2, column=6).value == "12345, 67890"  # Central Numbers (comma-separated)
    assert ws.cell(row=2, column=7).value == "ABC123, DEF456"  # Commitment Numbers


def test_excel_exporter_handles_empty_lists(tmp_path):
    """Test that empty lists are exported correctly."""
    record = TestRecord(
        test_id="test_002",
        document_id="doc_002",
        page_indices=[0],
        patient_first_name=FieldValue[str](
            raw_value="דני",
            extraction_state=ExtractionState.found
        ),
        patient_last_name=FieldValue[str](
            raw_value="לוי",
            extraction_state=ExtractionState.found
        ),
        patient_id_number=FieldValue[str](
            raw_value="987654321",
            extraction_state=ExtractionState.found
        ),
        sender=FieldValue[str](
            raw_value="רמב״ם",
            extraction_state=ExtractionState.found
        ),
        received_date=FieldValue[str](
            raw_value="2024-02-20",
            extraction_state=ExtractionState.found
        ),
        central_numbers=FieldValue[list[str]](
            raw_value=[],  # Empty list
            extraction_state=ExtractionState.missing
        ),
        commitment_ids=[]  # No commitments
    )
    
    exporter = ExcelExporter()
    output_path = tmp_path / "test_empty.xlsx"
    
    exporter.export(
        test_records=[record],
        commitment_records=[],
        output_path=output_path
    )
    
    wb = load_workbook(output_path)
    ws = wb.active
    
    assert ws.cell(row=2, column=6).value is None or ws.cell(row=2, column=6).value == ""  # Empty central numbers
    assert ws.cell(row=2, column=7).value is None or ws.cell(row=2, column=7).value == ""  # Empty commitments


def test_excel_exporter_handles_missing_values(tmp_path):
    """Test that None/missing values are handled correctly."""
    record = TestRecord(
        test_id="test_003",
        document_id="doc_003",
        page_indices=[0],
        patient_first_name=FieldValue[str](
            raw_value=None,  # Missing
            extraction_state=ExtractionState.missing
        ),
        patient_last_name=FieldValue[str](
            raw_value="ישראלי",
            extraction_state=ExtractionState.found
        ),
        patient_id_number=FieldValue[str](
            raw_value=None,  # Missing
            extraction_state=ExtractionState.missing
        ),
        sender=FieldValue[str](
            raw_value="איכילוב",
            extraction_state=ExtractionState.found
        ),
        received_date=FieldValue[str](
            raw_value=None,  # Missing
            extraction_state=ExtractionState.missing
        ),
        central_numbers=FieldValue[list[str]](
            raw_value=["11111"],
            extraction_state=ExtractionState.found
        ),
        commitment_ids=[]
    )
    
    exporter = ExcelExporter()
    output_path = tmp_path / "test_missing.xlsx"
    
    exporter.export(
        test_records=[record],
        commitment_records=[],
        output_path=output_path
    )
    
    wb = load_workbook(output_path)
    ws = wb.active
    
    assert ws.cell(row=2, column=1).value is None  # Missing first name
    assert ws.cell(row=2, column=3).value is None  # Missing ID
    assert ws.cell(row=2, column=5).value is None  # Missing date


def test_excel_exporter_invalid_output_path():
    """Test that exporter raises error for invalid output path."""
    exporter = ExcelExporter()
    
    # Non-existent parent directory
    invalid_path = Path("/nonexistent/directory/output.xlsx")
    
    with pytest.raises(ValueError, match="Output directory does not exist"):
        exporter.export(
            test_records=[],
            commitment_records=[],
            output_path=invalid_path
        )


def test_excel_exporter_uses_reviewed_values(tmp_path):
    """Test that exporter prioritizes reviewed > normalized > raw values."""
    record = TestRecord(
        test_id="test_004",
        document_id="doc_004",
        page_indices=[0],
        patient_first_name=FieldValue[str](
            raw_value="יוסף",
            normalized_value="יוסף",
            reviewed_value="יוסי",  # Reviewed value should be used
            extraction_state=ExtractionState.found,
            review_state=ReviewState.human_verified
        ),
        patient_last_name=FieldValue[str](
            raw_value="כהן",
            extraction_state=ExtractionState.found
        ),
        patient_id_number=FieldValue[str](
            raw_value="123456782",
            extraction_state=ExtractionState.found
        ),
        sender=FieldValue[str](
            raw_value="שיבא",
            extraction_state=ExtractionState.found
        ),
        received_date=FieldValue[str](
            raw_value="2024-01-15",
            extraction_state=ExtractionState.found
        ),
        central_numbers=FieldValue[list[str]](
            raw_value=["12345"],
            extraction_state=ExtractionState.found
        ),
        commitment_ids=[]
    )
    
    exporter = ExcelExporter()
    output_path = tmp_path / "test_reviewed.xlsx"
    
    exporter.export(
        test_records=[record],
        commitment_records=[],
        output_path=output_path
    )
    
    wb = load_workbook(output_path)
    ws = wb.active
    
    # Should use reviewed value "יוסי", not raw "יוסף"
    assert ws.cell(row=2, column=1).value == "יוסי"
