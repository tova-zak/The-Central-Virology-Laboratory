"""Offline end-to-end evidence: every accepted commitment number of a page reaches the Excel file.

One page with 3 commitment numbers and 2 central numbers goes through the real classifier,
the record builder, JSON persistence, GET, the verify endpoint and POST export; the actual
.xlsx is read back. No network, no Supabase, no Gemini.
"""
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import process
from app.services.extraction.commitment_classifier import classify_commitments
from app.services.extraction.extractor import ExtractionResult
from app.services.processing.background import BackgroundProcessor
from app.services.storage.json_storage import JSONStorage

COMMITMENT = "\u05d4\u05ea\u05d7\u05d9\u05d9\u05d1\u05d5\u05ea"
REQUEST = "\u05d3\u05e8\u05d9\u05e9\u05d4"
TEST_LABEL = "\u05d1\u05d3\u05d9\u05e7\u05d4"
BATCH = "batch_preserve"
REVIEW_FIELDS = ("patient_first_name", "patient_last_name", "patient_id_number",
                 "sender", "received_date", "central_numbers", "commitment_number")


def _page(extra_types=()):
    numbers = ["730578626", "730578627", "730578628"]
    types = [{"number": numbers[0], "label": COMMITMENT},
             {"number": numbers[1], "label": REQUEST},
             {"number": numbers[2], "label": COMMITMENT}, *extra_types]
    central = ["946621", "946622"]
    classification = classify_commitments(numbers + [t["number"] for t in extra_types], types, central, "Clalit")
    fields = {
        "patient_first_name": "Dana", "patient_last_name": "Levi", "patient_id_number": "012345674",
        "sender": "Clalit", "received_date": "15/03/2024", "central_numbers": central,
        "commitment_numbers": classification.accepted, "commitment_types": types,
        "commitment_candidates": classification.candidates,
    }
    result = ExtractionResult("1", fields, [], "synthetic", "test", 0.0)
    processor = object.__new__(BackgroundProcessor)
    return processor._create_commitment_from_page(result, "doc", 0)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(process, "_storage", JSONStorage(tmp_path))
    monkeypatch.setattr(process.settings, "temp_storage_path", str(tmp_path))
    app = FastAPI()
    app.include_router(process.router)
    return TestClient(app)


def _store(record):
    process._storage.save_batch(BATCH, {"test_records": [], "commitment_records": [record.model_dump(mode="json")]})


def _approve_and_export(client, record_id, tmp_path):
    for name in REVIEW_FIELDS:
        assert client.post(f"/api/batches/{BATCH}/tests/{record_id}/fields/{name}/verify").status_code == 200
    response = client.post(f"/api/batches/{BATCH}/export")
    assert response.status_code == 200
    out = tmp_path / "read_back.xlsx"
    out.write_bytes(response.content)
    return list(load_workbook(out).active.iter_rows(min_row=2, values_only=True))


def test_three_commitment_numbers_two_centrals_reach_excel(client, tmp_path):
    record = _page()
    _store(record)

    stored = client.get(f"/api/batches/{BATCH}").json()["commitment_records"][0]
    field = stored["commitment_number"]
    assert field["raw_value"] == "730578626"
    assert {c["value"] for c in field["candidates"]} == {"730578626", "730578627", "730578628"}
    assert stored["central_numbers"]["raw_value"] == ["946621", "946622"]

    rows = _approve_and_export(client, record.commitment_record_id, tmp_path)
    assert [r[6] for r in rows] == ["730578626", "730578627", "730578628"]
    assert all(r[5] == "946621, 946622" for r in rows)
    assert all(r[0:5] == rows[0][0:5] for r in rows)


def test_unlabelled_number_is_kept_as_candidate_but_not_exported(client, tmp_path):
    record = _page(extra_types=[{"number": "555", "label": TEST_LABEL}])
    _store(record)

    field = client.get(f"/api/batches/{BATCH}").json()["commitment_records"][0]["commitment_number"]
    kept = {c["value"]: c["reason"] for c in field["candidates"]}
    assert kept["555"] == "no_recognized_label"

    rows = _approve_and_export(client, record.commitment_record_id, tmp_path)
    assert [r[6] for r in rows] == ["730578626", "730578627", "730578628"]
