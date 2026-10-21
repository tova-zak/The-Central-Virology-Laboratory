"""Export integrity tests (offline, no network, no Supabase).

Covers:
- one Excel row per commitment number of a page (records are never merged)
- every exported cell is text, never a formula; leading zeros preserved
- confirmed-empty fields no longer block export (except patient_id_number / central_numbers)
- POST /api/batches/{id}/export end-to-end and PATCH rejecting the misread candidate
"""
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import process
from app.domain.models import (
    CommitmentRecord, FieldCandidate, FieldValue, ExtractionState, ReviewState,
)
from app.services.export.excel_exporter import ExcelExporter, commitment_numbers_for_export
from app.services.storage import JSONStorage

EXTRA = "multiple_commitments_on_page"


def _field(value, verified=True, candidates=None):
    return FieldValue(
        raw_value=value,
        normalized_value=value,
        extraction_state=ExtractionState.found,
        review_state=ReviewState.human_verified if verified else ReviewState.pending,
        candidates=candidates or [],
    )


def _record(record_id="c1", commitment="730578626", extra=(), verified=True, **overrides):
    candidates = [FieldCandidate(value=v, reason=EXTRA, status=s) for v, s in extra]
    fields = dict(
        patient_first_name=_field("Dana", verified),
        patient_last_name=_field("Levi", verified),
        patient_id_number=_field("012345674", verified),
        sender=_field("Clalit", verified),
        received_date=_field("15/03/2024", verified),
        central_numbers=_field(["946621", "946622", "946623"], verified),
        commitment_number=_field(commitment, verified, candidates),
        status="complete",
        relationship_status="inferred",
    )
    fields.update(overrides)
    return CommitmentRecord(commitment_record_id=record_id, form_id=f"{record_id}_page_1", **fields)


def _export(records, tmp_path):
    path = ExcelExporter().export([], records, tmp_path / "out.xlsx")
    return list(load_workbook(path).active.iter_rows(min_row=2))


def _gate(record):
    return process._check_export_eligibility_commitments([record.model_dump(mode="json")])


# (a) one row per commitment number --------------------------------------------------------------

def test_two_commitment_numbers_export_two_rows(tmp_path):
    rows = _export([_record(extra=[("730578627", "pending")])], tmp_path)
    assert [r[6].value for r in rows] == ["730578626", "730578627"]
    for column in range(6):
        assert rows[0][column].value == rows[1][column].value
    assert rows[0][5].value == "946621, 946622, 946623"


def test_rejected_candidate_is_not_exported(tmp_path):
    rows = _export([_record(extra=[("730578627", "rejected"), ("730578628", "pending")])], tmp_path)
    assert [r[6].value for r in rows] == ["730578626", "730578628"]


def test_other_candidate_reasons_and_duplicates_are_not_exported():
    record = _record(extra=[("730578626", "pending"), ("730578627", "pending")])
    record.commitment_number.candidates.append(FieldCandidate(value="111", reason="ambiguous_label"))
    assert commitment_numbers_for_export(record) == ["730578626", "730578627"]


def test_records_are_not_merged_and_row_counter_is_independent(tmp_path):
    records = [
        _record("c1", "100000001", extra=[("100000002", "pending")]),
        _record("c2", "200000001"),
    ]
    rows = _export(records, tmp_path)
    assert [r[6].value for r in rows] == ["100000001", "100000002", "200000001"]
    assert len(records) == 2


def test_empty_commitment_number_exports_one_empty_row(tmp_path):
    record = _record(commitment="")
    record.commitment_number.reviewed_value = ""
    record.commitment_number.candidates.append(FieldCandidate(value="9", reason=EXTRA))
    assert commitment_numbers_for_export(record) == [""]
    rows = _export([record], tmp_path)
    assert len(rows) == 1 and rows[0][6].value is None


# (b) text only ----------------------------------------------------------------------------------

@pytest.mark.parametrize("text", ['=HYPERLINK("http://x","a")', "+1+1", "-2+3", "@SUM(1)", "\t=1", "=1+1"])
def test_formula_like_text_is_stored_as_text(tmp_path, text):
    record = _record(patient_first_name=_field(text))
    cell = _export([record], tmp_path)[0][0]
    assert cell.data_type == "s"
    assert cell.value == text


def test_no_cell_is_a_formula_and_leading_zeros_survive(tmp_path):
    record = _record(
        commitment="007305786",
        patient_id_number=_field("012345674"),
        central_numbers=_field(["000123", "00456"]),
        sender=_field("=cmd|' /C calc'!A0"),
    )
    rows = _export([record], tmp_path)
    assert all(cell.data_type != "f" for row in rows for cell in row)
    assert rows[0][2].value == "012345674"
    assert rows[0][5].value == "000123, 00456"
    assert rows[0][6].value == "007305786"
    assert rows[0][3].value == "=cmd|' /C calc'!A0"


# (c) confirmed-empty gate -----------------------------------------------------------------------

@pytest.mark.parametrize("name", ["commitment_number", "sender", "patient_first_name",
                                  "patient_last_name", "received_date"])
def test_confirmed_empty_field_does_not_block(name):
    empty = FieldValue(reviewed_value="", review_state=ReviewState.human_verified)
    ok, blockers = _gate(_record(**{name: empty}))
    assert ok, blockers


@pytest.mark.parametrize("name,empty_value", [("patient_id_number", ""), ("central_numbers", [])])
def test_required_fields_cannot_be_confirmed_empty(name, empty_value):
    empty = FieldValue(reviewed_value=empty_value, review_state=ReviewState.human_verified)
    ok, blockers = _gate(_record(**{name: empty}))
    assert not ok and f"c1: {name}_missing" in blockers


@pytest.mark.parametrize("name", ["commitment_number", "sender", "patient_id_number"])
def test_unreviewed_empty_field_still_blocks(name):
    empty = FieldValue(review_state=ReviewState.pending)
    ok, blockers = _gate(_record(**{name: empty}))
    assert not ok and f"c1: {name}_missing" in blockers


def test_unreviewed_non_empty_field_still_blocks():
    ok, blockers = _gate(_record(sender=_field("Clalit", verified=False)))
    assert not ok and "c1: sender_unreviewed" in blockers


# (d) endpoint -----------------------------------------------------------------------------------

@pytest.fixture
def client(tmp_path, monkeypatch):
    from app.main import app
    storage = JSONStorage(base_path=tmp_path)
    monkeypatch.setattr(process, "_storage", storage)
    monkeypatch.setattr(process.settings, "temp_storage_path", str(tmp_path))
    return TestClient(app), storage


def _save(storage, record, batch_id="batch_integrity"):
    storage.save_batch(batch_id, {
        "batch_id": batch_id,
        "status": "completed",
        "test_records": [],
        "commitment_records": [record.model_dump(mode="json")],
    })
    return batch_id


def test_export_endpoint_returns_one_row_per_commitment(client, tmp_path):
    http, storage = client
    batch_id = _save(storage, _record(extra=[("730578627", "pending")]))
    response = http.post(f"/api/batches/{batch_id}/export")
    assert response.status_code == 200
    out = tmp_path / "downloaded.xlsx"
    out.write_bytes(response.content)
    rows = list(load_workbook(out).active.iter_rows(min_row=2, values_only=True))
    assert [r[6] for r in rows] == ["730578626", "730578627"]
    assert len(storage.load_batch(batch_id)["commitment_records"]) == 1


def test_export_endpoint_still_blocks_unreviewed_field(client):
    http, storage = client
    batch_id = _save(storage, _record(sender=_field("Clalit", verified=False)))
    response = http.post(f"/api/batches/{batch_id}/export")
    assert response.status_code == 400
    assert "c1: sender_unreviewed" in response.json()["detail"]["blockers"]


def test_correcting_commitment_number_rejects_misread_candidate(client):
    http, storage = client
    record = _record(
        verified=False,
        extra=[("730578626", "pending"), ("730578627", "pending")],
    )
    batch_id = _save(storage, record)
    url = f"/api/batches/{batch_id}/tests/c1/fields/commitment_number"
    response = http.patch(url, json={"new_value": "730578999"})
    assert response.status_code == 200
    field = storage.load_batch(batch_id)["commitment_records"][0]["commitment_number"]
    status = {c["value"]: c["status"] for c in field["candidates"]}
    assert status == {"730578626": "rejected", "730578627": "pending"}
    assert field["reviewed_value"] == "730578999"
    assert field["review_state"] == "human_verified"
