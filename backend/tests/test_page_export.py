"""Offline regression coverage for page-owned patient fields and export approval."""
from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import process
from app.domain.models import CommitmentRecord
from app.services.extraction.mock_extractor import MockExtractor
from app.services.processing.background import BackgroundProcessor
from app.services.storage.json_storage import JSONStorage


@pytest.fixture
def page_batch(tmp_path, monkeypatch):
    monkeypatch.setattr(process, '_storage', JSONStorage(tmp_path))
    monkeypatch.setattr(process.settings, 'temp_storage_path', str(tmp_path))
    # Avoid initializing any live extractor or background worker.
    processor = object.__new__(BackgroundProcessor)
    extractor = MockExtractor(simulate_delay=False)
    records = []
    for index in range(3):
        result = extractor.extract(b'', 'synthetic_document', index)
        record = processor._create_commitment_from_page(result, 'synthetic_document', index)
        assert record.patient_first_name.raw_value == result.fields['patient_first_name']
        assert record.central_numbers.raw_value == result.fields['central_numbers']
        records.append(record.model_dump(mode='json'))
    process._storage.save_batch('synthetic_batch', {'test_records': [], 'commitment_records': records})
    app = FastAPI()
    app.include_router(process.router)
    return TestClient(app)


def approve(client):
    data = process._storage.load_batch('synthetic_batch')
    for record in data['commitment_records']:
        for name in ('patient_first_name', 'patient_last_name', 'patient_id_number',
                     'sender', 'received_date', 'central_numbers', 'commitment_number'):
            response = client.post(f"/api/batches/synthetic_batch/tests/{record['commitment_record_id']}/fields/{name}/verify")
            assert response.status_code == 200


def test_page_fields_survive_storage_review_and_post_export(page_batch):
    client = page_batch
    assert client.post('/api/batches/synthetic_batch/export').status_code == 400
    approve(client)
    record = process._storage.load_batch('synthetic_batch')['commitment_records'][0]
    path = f"/api/batches/synthetic_batch/tests/{record['commitment_record_id']}/fields/patient_first_name"
    assert client.patch(path, json={'new_value': 'Corrected synthetic name'}).status_code == 200
    assert client.get('/api/batches/synthetic_batch/export-check').json()['can_export'] is True
    response = client.post('/api/batches/synthetic_batch/export')
    assert response.status_code == 200
    ws = load_workbook(BytesIO(response.content)).active
    stored = process._storage.load_batch('synthetic_batch')['commitment_records']
    # Export is one Excel row per accepted commitment number (P0 commitment preservation):
    # the primary number first, then every non-rejected 'multiple_commitments_on_page' candidate.
    # The synthetic page 0 carries a second accepted number, so 3 pages -> 4 data rows (+ header).
    expected = []
    for data in stored:
        record = CommitmentRecord(**data)
        numbers = [record.commitment_number.raw_value] + [
            c.value for c in record.commitment_number.candidates
            if c.reason == 'multiple_commitments_on_page' and c.status != 'rejected'
            and c.value != record.commitment_number.raw_value]
        for number in numbers:
            expected.append((record.patient_first_name.reviewed_value or record.patient_first_name.raw_value,
                             record.patient_last_name.raw_value, record.patient_id_number.raw_value,
                             record.sender.raw_value, record.received_date.raw_value,
                             ', '.join(record.central_numbers.raw_value), number))
    assert len(expected) == 4 and ws.max_column == 7
    assert ws.max_row == len(expected) + 1
    assert list(ws.iter_rows(min_row=2, values_only=True)) == expected
    assert ws.cell(2, 1).value == 'Corrected synthetic name'
    assert ws.cell(3, 1).value == 'Corrected synthetic name'  # second number of the same page


@pytest.mark.parametrize('name', ['patient_first_name', 'patient_last_name', 'patient_id_number',
                                 'sender', 'received_date', 'central_numbers', 'commitment_number'])
@pytest.mark.parametrize('problem', ['absent', 'empty_correction', 'empty_unreviewed', 'unreviewed'])
def test_required_field_blockers(page_batch, name, problem):
    approve(page_batch)
    data = process._storage.load_batch('synthetic_batch')
    record = data['commitment_records'][0]
    blocks = True
    if problem == 'absent':
        del record[name]
    elif problem == 'empty_correction':
        # A reviewer-approved (human_verified) empty value IS a confirmed-empty field (D-17 / P0-3):
        # it no longer blocks export, except for the fields policy requires (ID, central numbers).
        record[name]['reviewed_value'] = [] if name == 'central_numbers' else '   '
        blocks = name in ('patient_id_number', 'central_numbers')
    elif problem == 'empty_unreviewed':
        # Empty WITHOUT a human confirmation always blocks.
        record[name]['reviewed_value'] = [] if name == 'central_numbers' else '   '
        record[name]['review_state'] = 'needs_review'
    else:
        record[name]['review_state'] = 'needs_review'
    process._storage.save_batch('synthetic_batch', data)
    assert page_batch.get('/api/batches/synthetic_batch/export-check').json()['can_export'] is (not blocks)
    assert page_batch.post('/api/batches/synthetic_batch/export').status_code == (400 if blocks else 200)


def test_confirm_empty_via_edit_endpoint_roundtrip_and_export(page_batch, tmp_path):
    approve(page_batch)
    rid = process._storage.load_batch('synthetic_batch')['commitment_records'][1]['commitment_record_id']
    base = f"/api/batches/synthetic_batch/tests/{rid}/fields"
    assert page_batch.patch(f"{base}/sender", json={'new_value': ''}).status_code == 200
    # required-by-policy fields cannot be cleared
    assert page_batch.patch(f"{base}/patient_id_number", json={'new_value': ''}).status_code == 400
    assert page_batch.patch(f"{base}/central_numbers", json={'new_value': []}).status_code == 400
    # reload through a brand-new storage instance
    fresh = JSONStorage(tmp_path)
    field = next(r for r in fresh.load_batch('synthetic_batch')['commitment_records']
                 if r['commitment_record_id'] == rid)['sender']
    assert field['reviewed_value'] == '' and field['review_state'] == 'human_verified'
    assert field['raw_value']  # original extracted value still stored
    ws = load_workbook(BytesIO(page_batch.post('/api/batches/synthetic_batch/export').content)).active
    assert [r[3] for r in ws.iter_rows(min_row=2, values_only=True)].count(None) == 1


@pytest.mark.parametrize('text', ['=1+1', '+972501234567', '-5', '@SUM(A1)', '\t=1', '\r=1'])
def test_reviewed_text_roundtrip_and_formula_injection(page_batch, tmp_path, text):
    approve(page_batch)
    rid = process._storage.load_batch('synthetic_batch')['commitment_records'][0]['commitment_record_id']
    path = f"/api/batches/synthetic_batch/tests/{rid}/fields/sender"
    assert page_batch.patch(path, json={'new_value': text}).status_code == 200
    raw_before = process._storage.load_batch('synthetic_batch')['commitment_records'][0]['sender']['raw_value']
    reloaded = JSONStorage(tmp_path).load_batch('synthetic_batch')['commitment_records'][0]['sender']
    assert reloaded['reviewed_value'] == text and reloaded['raw_value'] == raw_before
    ws = load_workbook(BytesIO(page_batch.post('/api/batches/synthetic_batch/export').content)).active
    cell = ws.cell(2, 4)
    # Stored verbatim as text, never a formula. Only a lone CR is normalised to LF by the XLSX/XML layer.
    assert cell.data_type == 's' and cell.value == text.replace('\r', '\n')
    assert all(c.data_type != 'f' for row in ws.iter_rows() for c in row)
    # neighbouring identifier cells are untouched by the neutralisation
    id_raw = process._storage.load_batch('synthetic_batch')['commitment_records'][0]['patient_id_number']['raw_value']
    assert ws.cell(2, 3).value == id_raw


def test_empty_batch_and_unresolved_relationship_block(page_batch):
    approve(page_batch)
    data = process._storage.load_batch('synthetic_batch')
    data['commitment_records'][0]['relationship_status'] = 'unresolved_relationship'
    process._storage.save_batch('synthetic_batch', data)
    assert page_batch.post('/api/batches/synthetic_batch/export').status_code == 400
    assert process._check_export_eligibility_commitments([])[0] is False


def test_reviewed_empty_never_falls_back():
    from app.domain.models import FieldValue
    from app.services.export.excel_exporter import ExcelExporter
    assert ExcelExporter()._get_value(FieldValue[str](raw_value='original', reviewed_value='')) == ''
